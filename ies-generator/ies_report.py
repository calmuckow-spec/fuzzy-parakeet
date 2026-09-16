# -*- coding: utf-8 -*-
"""
ies_report.py — выгрузка сводок по файлам IES в Excel.
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT_NAME = "Arial"

_HEAD_FILL = PatternFill("solid", fgColor="D9D9D9")
_INPUT_FILL = PatternFill("solid", fgColor="FFFF99")
_THIN = Side(style="thin", color="999999")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _write_table(ws, headers, rows, widths, formulas=None, start_row=1,
                 number_formats=None):
    """
    headers  — список заголовков;
    rows     — список списков значений (в местах формул ставится None);
    formulas — словарь {номер столбца: шаблон формулы с {r}};
    number_formats — словарь {номер столбца: формат числа}.
    Возвращает номер последней заполненной строки таблицы.
    """
    head_row = start_row
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=head_row, column=col, value=text)
        cell.font = Font(name=FONT_NAME, bold=True)
        cell.fill = _HEAD_FILL
        cell.border = _BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r_off, row in enumerate(rows):
        r = head_row + 1 + r_off
        for c_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c_idx, value=value)
            cell.font = Font(name=FONT_NAME)
            cell.border = _BORDER
            if isinstance(value, (int, float)):
                cell.alignment = Alignment(horizontal="right")
        if formulas:
            for col, tpl in formulas.items():
                cell = ws.cell(row=r, column=col, value=tpl.format(r=r))
                cell.font = Font(name=FONT_NAME)
                cell.border = _BORDER
                cell.alignment = Alignment(horizontal="right")
        if number_formats:
            for col, fmt in number_formats.items():
                ws.cell(row=r, column=col).number_format = fmt

    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.freeze_panes = ws.cell(row=head_row + 1, column=1).coordinate
    return head_row + len(rows)


def _note(ws, row, text):
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(name=FONT_NAME, italic=True, size=9)
    return row


# --------------------------------------------------------- отчёт 1: созданные

def export_created(path, items):
    """
    Сводка по созданным файлам.
    items — словари: name, flux, watts, length, width, imax.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Созданные файлы"

    headers = [
        "Название",
        "Световой поток, лм",
        "Мощность, Вт",
        "Световая отдача, лм/Вт",
        "Длина, м",
        "Ширина, м",
        "Макс. осевая сила света, кд",
        "Габаритная яркость, кд/м²",
    ]
    rows = []
    for it in items:
        rows.append([
            it["name"], it["flux"], it["watts"], None,
            it["length"], it["width"], it["imax"], None,
        ])

    last = _write_table(
        ws, headers, rows,
        widths=[42, 20, 15, 20, 11, 11, 26, 24],
        formulas={
            4: '=IFERROR(B{r}/C{r},"")',
            8: '=IFERROR(G{r}/(E{r}*F{r}),"")',
        },
        number_formats={4: "0.0", 5: "0.000", 6: "0.000", 7: "0.0", 8: "#,##0"},
    )

    _note(ws, last + 2,
          "Габаритная яркость = максимальная осевая сила света / (длина × ширина светового окна). "
          "Размеры взяты из параметров, заданных при создании файлов.")
    wb.save(path)
    return path


# ------------------------------------------------------------ отчёт 2: папка

def export_scan(path, items):
    """
    Сводка по файлам IES из выбранной папки.
    items — словари: name, flux, watts, length, width, height, imax, zonal.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Анализ папки"

    headers = [
        "Название",
        "Световой поток, лм",
        "Мощность, Вт",
        "Световая отдача, лм/Вт",
        "Длина, м",
        "Ширина, м",
        "Высота, м",
        "Макс. осевая сила света, кд",
        "Габаритная яркость, кд/м²",
        "Поток по кривой силы света, лм",
    ]
    rows = []
    for it in items:
        rows.append([
            it["name"], it["flux"], it["watts"], None,
            it["length"], it["width"], it["height"],
            it["imax"], None, it["zonal"],
        ])

    last = _write_table(
        ws, headers, rows,
        widths=[42, 20, 15, 20, 11, 11, 11, 26, 24, 24],
        formulas={
            4: '=IFERROR(B{r}/C{r},"")',
            9: '=IFERROR(H{r}/(E{r}*F{r}),"")',
        },
        number_formats={4: "0.0", 5: "0.000", 6: "0.000", 7: "0.000",
                        8: "0.0", 9: "#,##0", 10: "0.0"},
    )

    _note(ws, last + 2,
          "Габаритная яркость = максимальная осевая сила света / (длина × ширина светового окна). "
          "Размеры взяты из самого файла IES.")
    _note(ws, last + 3,
          "Световой поток взят из поля «люмены на лампу». Последний столбец — "
          "независимая проверка: поток, посчитанный интегрированием кривой силы света.")
    wb.save(path)
    return path


# ------------------------------------------------- отчёт 3: габаритная яркость

def export_luminance(path, items, win_length, win_width):
    """
    Расчёт габаритной яркости по заданному размеру светового окна.
    items — словари: name, flux, watts, iaxial.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Габаритная яркость"

    ws["A1"] = "Размер светового окна"
    ws["A1"].font = Font(name=FONT_NAME, bold=True)

    ws["A2"] = "Длина, м"
    ws["B2"] = win_length
    ws["A3"] = "Ширина, м"
    ws["B3"] = win_width
    ws["A4"] = "Площадь, м²"
    ws["B4"] = "=B2*B3"

    for row in (2, 3, 4):
        ws.cell(row=row, column=1).font = Font(name=FONT_NAME)
        cell = ws.cell(row=row, column=2)
        cell.font = Font(name=FONT_NAME)
        cell.border = _BORDER
        cell.number_format = "0.0000"
        if row < 4:
            cell.fill = _INPUT_FILL

    _note(ws, 5, "Жёлтые ячейки — исходные данные, их можно менять: "
                 "яркость в таблице пересчитается сама.")

    headers = [
        "Название светового прибора",
        "Световой поток, лм",
        "Мощность, Вт",
        "Световая отдача, лм/Вт",
        "Макс. осевая сила света, кд",
        "Габаритная яркость, кд/м²",
    ]
    rows = [[it["name"], it["flux"], it["watts"], None, it["iaxial"], None]
            for it in items]

    last = _write_table(
        ws, headers, rows,
        widths=[46, 20, 15, 20, 26, 26],
        formulas={
            4: '=IFERROR(B{r}/C{r},"")',
            6: '=IFERROR(E{r}/$B$4,"")',
        },
        number_formats={4: "0.0", 5: "0.0", 6: "#,##0"},
        start_row=7,
    )

    _note(ws, last + 2,
          "Габаритная яркость = максимальная осевая сила света / площадь светового окна.")
    wb.save(path)
    return path
